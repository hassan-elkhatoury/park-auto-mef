from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPORT_BLUE = "1F487C"
REPORT_LIGHT = "F3F6FA"
REPORT_BORDER = "B8C9DC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str, size: int = 8) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_twips: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    grid_col = OxmlElement("w:gridCol")
    grid_col.set(qn("w:w"), str(width_twips))
    grid.append(grid_col)
    set_cell_width(table.cell(0, 0), width_twips)


def add_roadmap(document: Document) -> None:
    intro = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("Ce chapitre présente l'ensemble des travaux de conception")
    )
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 9220)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, REPORT_LIGHT)
    set_cell_borders(cell, REPORT_BLUE, size=8)

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.05

    label = paragraph.add_run("Organisation du chapitre. ")
    label.bold = True
    label.font.name = "Aptos"
    label.font.size = Pt(10.5)
    label.font.color.rgb = RGBColor(11, 37, 69)

    body = paragraph.add_run(
        "La conception est présentée selon six blocs cohérents : socle technique et architecture, "
        "sécurité et habilitations, gestion du parc, conducteurs et missions, exploitation "
        "carburant/maintenance, puis risques et pilotage budgétaire. Cette progression relie les "
        "besoins fonctionnels aux modèles UML, aux données et aux interfaces React."
    )
    body.font.name = "Aptos"
    body.font.size = Pt(10.5)
    body.font.color.rgb = RGBColor(32, 39, 51)

    intro._p.addprevious(table._tbl)


def set_image_alt_text(inline_shape, description: str) -> None:
    inline_shape._inline.docPr.set("descr", description)
    inline_shape._inline.docPr.set("title", description)
    c_nv_pr_nodes = inline_shape._inline.xpath(".//pic:cNvPr")
    if c_nv_pr_nodes:
        c_nv_pr_nodes[0].set("descr", description)


def insert_figure(document: Document, caption_prefix: str, image_path: Path, width_inches: float, alt_text: str) -> None:
    paragraphs = document.paragraphs
    caption_index = next(
        index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip().startswith(caption_prefix)
    )
    image_paragraph = paragraphs[caption_index - 1]
    image_paragraph.clear()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(2)
    image_paragraph.paragraph_format.space_after = Pt(4)
    image_paragraph.paragraph_format.keep_with_next = True
    run = image_paragraph.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(width_inches))
    set_image_alt_text(inline_shape, alt_text)

    caption = paragraphs[caption_index]
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(6)


def enable_field_updates(document: Document) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def format_chapter3_headings(document: Document) -> None:
    chapter_started = False
    full_page_figure_sections = {
        "3.2.2. Cas d'utilisation",
        "3.2.3. Modélisation statique",
        "3.2.4. Modélisation dynamique",
    }
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Chapitre 3 - Architecture, conception et réalisation"):
            chapter_started = True
        if not chapter_started:
            continue
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.widow_control = True
        if text in full_page_figure_sections:
            paragraph.paragraph_format.page_break_before = True


def build(source: Path, output: Path, assets: Path) -> None:
    if source.resolve() == output.resolve():
        raise ValueError("The output path must differ from the source DOCX.")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)

    document = Document(output)
    add_roadmap(document)
    insert_figure(
        document,
        "Figure 13 -",
        assets / "CH3_Asset_05_UseCase_Securite.png",
        2.35,
        "Diagramme de cas d'utilisation des fonctions de sécurité et d'administration",
    )
    insert_figure(
        document,
        "Figure 14 -",
        assets / "CH3_Asset_06_Classes_Securite.png",
        5.25,
        "Diagramme de classes des habilitations et des services transverses",
    )
    insert_figure(
        document,
        "Figure 15 -",
        assets / "CH3_Asset_07_Sequence_Login_JWT.png",
        6.15,
        "Diagramme de séquence du contrôle d'accès JWT",
    )
    format_chapter3_headings(document)
    enable_field_updates(document)
    document.core_properties.title = "Rapport PFA - Gestion du Parc Automobile MEF"
    document.core_properties.subject = "Version Word générée depuis le projet LaTeX final"
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the final DOCX version from the LaTeX report project.")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("assets", type=Path)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve(), args.assets.resolve())


if __name__ == "__main__":
    main()
