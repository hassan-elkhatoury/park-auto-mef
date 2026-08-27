import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = docx.Document()

# Margins
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

# Colors matching MEF style
NAVY = RGBColor(15, 29, 50)          # #0F1D32
BLUE_ACCENT = RGBColor(79, 129, 189) # #4F81BD
DARK_GRAY = RGBColor(60, 60, 60)

# Title
p_t = doc.add_paragraph()
r_t1 = p_t.add_run('Planning du Sprint 6 (Conforme CdC MEF & Recommandations Agile)\n')
r_t1.bold = True
r_t1.font.size = Pt(18)
r_t1.font.color.rgb = NAVY

r_t2 = p_t.add_run('Gestion de la Maintenance, des Pannes & Traitement des Sinistres (17/08 – 21/08/2026)')
r_t2.bold = True
r_t2.font.size = Pt(11.5)
r_t2.font.color.rgb = BLUE_ACCENT

p_t.paragraph_format.space_after = Pt(14)

# 1. Informations Générales
h1 = doc.add_heading('1. Informations Générales du Sprint 6', level=1)
h1.runs[0].font.color.rgb = NAVY
h1.runs[0].font.size = Pt(12.5)

table_info = doc.add_table(rows=4, cols=2)
table_info.alignment = WD_TABLE_ALIGNMENT.CENTER
table_info.autofit = False

info_data = [
    ('Sprint', 'Sprint 6 — Maintenance, Pannes & Sinistres'),
    ('Période prévisionnelle', '17/08/2026 au 21/08/2026 (5 jours ouvrés)'),
    ('Conformité CdC MEF', 'Sections 14 (Sinistres), 15 (Entretiens) & 16 (Pannes et Réparations) du Cahier des Charges.'),
    ('Objectif Principal', 'Spécifier, concevoir et développer le système global de suivi des entretiens préventifs, des interventions curatives sur pannes, de la gestion des garages agréés, du remplacement des pièces détachées et du processus complet de traitement et suivi des sinistres/accidents du parc automobile MEF.')
]

for idx, (label, val) in enumerate(info_data):
    r_cells = table_info.rows[idx].cells
    r_cells[0].text = label
    r_cells[0].paragraphs[0].runs[0].font.bold = True
    r_cells[0].paragraphs[0].runs[0].font.color.rgb = NAVY
    r_cells[1].text = val

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# 2. Contenu Détaillé du Sprint 6
h2 = doc.add_heading('2. Contenu Détaillé du Sprint 6', level=1)
h2.runs[0].font.color.rgb = NAVY
h2.runs[0].font.size = Pt(12.5)

h2_1 = doc.add_heading('2.1. Backlog Fonctionnel (Exigences CdC MEF)', level=2)
h2_1.runs[0].font.color.rgb = BLUE_ACCENT
h2_1.runs[0].font.size = Pt(11)

features = [
    ('1. Maintenance Préventive & Planification des Entretiens (CdC Section 15)', [
        'Programmation des révisions périodiques par seuil kilométrique (ex: 10 000 km, 20 000 km) ou fréquence temporelle.',
        'Déclenchement automatique d\'alerte de maintenance préventive lorsque le véhicule atteint 90% du seuil kilométrique.',
        'Gestion des types d\'entretiens : vidange, remplacement de filtres, contrôle des freins, remplacement pneumatiques, entretien batterie, climatisation et révision générale.',
        'Enregistrement de la réalisation : date réelle, kilométrage effectif, prestataire/garage agréé, détail des pièces remplacées, coût main d\'œuvre, coût pièces, facture et calcul de la prochaine échéance.'
    ]),
    ('2. Gestion des Pannes & Réparations Curatives (CdC Section 16)', [
        'Déclaration de panne : enregistrement des pannes par le conducteur/gestionnaire (véhicule, conducteur, date, lieu, kilométrage, nature de la panne, urgence, possibilité de déplacement/remorquage).',
        'Diagnostic atelier & Devis : enregistrement du diagnostic du garage agréé, origines de la panne, durée d\'immobilisation prévue, coût estimé et avis technique.',
        'Exécution & Clôture de réparation : suivi des travaux réalisés, des pièces détachées remplacées, bon de sortie d\'atelier, coût réel TTC, attachement de la facture et garantie accordée.',
        'Mise à jour automatique du statut du véhicule (passage à EN_MAINTENANCE / EN_REPARATION puis retour automatique à DISPONIBLE lors de la clôture).'
    ]),
    ('3. Traitement et Suivi Intégral des Sinistres Automobile (CdC Section 14)', [
        'Déclaration de sinistre/accident : enregistrement immédiat (date, heure, lieu, véhicule, conducteur, circonstances, tiers impliqués, constat amiable/PV de police, photographies, estimation des dommages).',
        'Suivi du dossier auprès de la compagnie d\'assurance : gestion du workflow des statuts (DECLARE, TRANSMIS, EN_COURS_D_EXPERTISE, ACCEPTE, REJETE, INDEMNISE, CLOTURE).',
        'Gestion financière du sinistre : suivi des montants des expertises, devis de réparation, montant de la franchise restée à charge, montant remboursé par l\'assurance et clôture finale.',
        'Bascule automatique du véhicule au statut ACCIDENTE ou EN_REPARATION avec blocage des nouvelles affectations tant que le sinistre n\'est pas clôturé.'
    ]),
    ('4. Intégration Budgétaire, Alertes & Gestion Documentaire (CdC Section 19, 21 & 22)', [
        'Impact budgétaire automatique : comptabilisation des dépenses d\'entretien, de réparation et de sinistre dans le budget alloué à la Direction/Service concerné.',
        'Alertes multi-canaux (Dashboard In-App & Email SMTP via JavaMailSender) pour dépassement de délais d\'immobilisation, retard d\'expertise ou coût de réparation anormal.',
        'Attachement systématique des pièces justificatives via le module GED (devis, factures, PV de constat, rapports d\'expertise, photos de dommages).'
    ])
]

for title, points in features:
    p_feat = doc.add_paragraph()
    r_f = p_feat.add_run(f'• {title}\n')
    r_f.bold = True
    r_f.font.color.rgb = NAVY
    for pt in points:
        p_pt = doc.add_paragraph()
        p_pt.paragraph_format.left_indent = Inches(0.2)
        r_bullet = p_pt.add_run('- ')
        r_bullet.bold = True
        p_pt.add_run(pt)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# 2.2 Priorisation MoSCoW & Règles de Gestion
h2_2 = doc.add_heading('2.2. Priorisation MoSCoW & Règles de Gestion Métier', level=2)
h2_2.runs[0].font.color.rgb = BLUE_ACCENT
h2_2.runs[0].font.size = Pt(11)

moscow_text = [
    'MoSCoW - Must Have (Priorité 1 — Impératif) : Déclaration/Clôture des Pannes, Planification Préventive (seuil kilométrique 90%), Gestion des Sinistres & Bascule automatique des statuts (EN_REPARATION, ACCIDENTE, DISPONIBLE).',
    'MoSCoW - Should Have (Priorité 2 — Essentiel) : Répertoire centralisé des Garages Agréés, Saisie détaillée du catalogue de pièces détachées et gestion des factures via GED.',
    'MoSCoW - Could Have (Priorité 3 — Optionnel) : Ordre de réparation automatique déclenché sur détection de contre-visite technique sous 15 jours.'
]

for m in moscow_text:
    p_m = doc.add_paragraph()
    p_m.paragraph_format.left_indent = Inches(0.1)
    r_mb = p_m.add_run('📌 ')
    r_mb.font.color.rgb = NAVY
    p_m.add_run(m)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

rules = [
    'RG01 — Un déclenchement d\'entretien préventif s\'active automatiquement dès que le kilométrage atteint 90% du seuil fixé (ex: 9 000 km pour un seuil à 10 000 km).',
    'RG02 — La création d\'une intervention de maintenance lourde ou la déclaration d\'une panne immobilisante passe immédiatement le véhicule au statut EN_MAINTENANCE ou EN_REPARATION et interdit toute nouvelle affectation.',
    'RG03 — Une déclaration de sinistre grave passe le véhicule au statut ACCIDENTE. La remise en circulation nécessite la clôture de l\'expertise et un PV de réparation.',
    'RG04 — Contrôle du Kilométrage Croissant : Le kilométrage réel de clôture ne peut être inférieur au kilométrage actuel du véhicule (if (kmReel < vehicule.getKilometrageActuel()) throw Exception).',
    'RG05 — Tout coût de réparation dépassant le budget restant alloué à la Direction du MEF requiert une alerte automatique et une validation hiérarchique.'
]

for rule in rules:
    p_r = doc.add_paragraph()
    p_r.paragraph_format.left_indent = Inches(0.1)
    r_b = p_r.add_run('✔ ')
    r_b.font.color.rgb = BLUE_ACCENT
    p_r.add_run(rule)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# 3. Recommandations Techniques & Plan d'Action Agile
h3_rec = doc.add_heading('3. Recommandations Techniques & Plan d\'Action Agile', level=1)
h3_rec.runs[0].font.color.rgb = NAVY
h3_rec.runs[0].font.size = Pt(12.5)

recs = [
    ('1. Schéma Relationnel Unique pour Sinistre et Intervention (@ManyToOne)', [
        'Validation de la modélisation JPA : S\'assurer que les entités Sinistre et InterventionMaintenance possèdent des relations directes @ManyToOne vers Vehicule, Conducteur et GarageAgree.',
        'Garantir l\'intégrité référentielle en BDD PostgreSQL et la traçabilité immédiate du conducteur au moment du sinistre ou de la panne.'
    ]),
    ('2. Validation Stricte du Kilométrage Croissant dans les Services Java (RG04)', [
        'Enforcement dans MaintenanceService et PanneService de la règle de non-rétrogradation :',
        'java code: if (kmReel < vehicule.getKilometrageActuel()) { throw new IllegalArgumentException("Le kilométrage de clôture ne peut pas être inférieur au kilométrage actuel du véhicule."); }',
        'Mise à jour atomique de vehicule.setKilometrageActuel(kmReel) lors de la clôture de l\'intervention.'
    ]),
    ('3. Application Rigoureuse de la Priorisation MoSCoW pendant le Codage', [
        'Must Have (P1) : Focus immédiat sur les pannes, entretiens (90% km), sinistres et la bascule automatique de statuts (EN_REPARATION, ACCIDENTE, DISPONIBLE).',
        'Should Have (P2) : Intégration du répertoire des garages agréés, pièces détachées et téléversement des factures dans la GED.'
    ])
]

for title, points in recs:
    p_rec = doc.add_paragraph()
    r_rc = p_rec.add_run(f'💡 {title}\n')
    r_rc.bold = True
    r_rc.font.color.rgb = NAVY
    for pt in points:
        p_pt = doc.add_paragraph()
        p_pt.paragraph_format.left_indent = Inches(0.2)
        r_bullet = p_pt.add_run('- ')
        r_bullet.bold = True
        p_pt.add_run(pt)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# 4. Livrables Attendus
h4 = doc.add_heading('4. Livrables Attendus du Sprint 6', level=1)
h4.runs[0].font.color.rgb = NAVY
h4.runs[0].font.size = Pt(12.5)

deliverables = [
    'Base de Données PostgreSQL : Entités JPA InterventionMaintenance, PlanificationEntretien, Panne, RepairOrder, Sinistre, GarageAgree, PieceRemplacement avec liaisons @ManyToOne vers Vehicule, Conducteur et GarageAgree.',
    'Services & Logic Métier Java : Services Spring Boot MaintenanceService, PanneService, SinistreService, GarageService avec contrôle de kilométrage croissant (kmReel >= kmActuel), alerte 90% km et notifications SMTP JavaMailSender.',
    'Endpoints API REST & Swagger UI : Controllers sécurisés /api/maintenance, /api/pannes, /api/reparations, /api/sinistres et /api/garages.',
    'Interface Frontend React : Vues /maintenance (planification & suivi entretiens), /pannes (déclarations & diagnostics), /sinistres (dossiers & expertises) et modales de clôture avec GED.',
    'Jeu de Données de Test & Documentation : Scripts d\'insertion de jeux de tests d\'entretiens/sinistres, guides utilisateurs et cahier de recette UAT des modules Maintenance & Sinistres.',
    'Démonstration & Validation : Vendredi 21 août 2026 à 11h00 au Ministère de l\'Économie et des Finances.'
]

for dev in deliverables:
    p_d = doc.add_paragraph()
    p_d.paragraph_format.left_indent = Inches(0.1)
    r_b = p_d.add_run('✔ ')
    r_b.font.color.rgb = NAVY
    p_d.add_run(dev)

output_docx_scripts = r'c:\Users\Hassan\Desktop\park auto MEF\scripts\Planning_Sprint_6_Maintenance_Pannes_Sinistres.docx'
output_docx_planning = r'c:\Users\Hassan\Desktop\park auto MEF\planning\Planning_Sprint_6_Maintenance_Pannes_Sinistres.docx'
output_docx_root = r'c:\Users\Hassan\Desktop\park auto MEF\Planning_Sprint_6_Maintenance_Pannes_Sinistres.docx'

doc.save(output_docx_scripts)
doc.save(output_docx_planning)
doc.save(output_docx_root)
print('DOCX Sprint 6 avec Recommandations Agile cree avec succes dans scripts, planning et root!')
