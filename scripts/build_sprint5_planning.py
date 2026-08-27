import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = docx.Document()

# Margins
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

# Colors matching Sprint 4 exactly
NAVY = RGBColor(7, 13, 27)
BLUE_ACCENT = RGBColor(79, 129, 189) # #4F81BD
DARK_GRAY = RGBColor(60, 60, 60)

# Title
p_t = doc.add_paragraph()
r_t1 = p_t.add_run('Planning du Sprint 5 (Version Finale — 100% CdC MEF)\n')
r_t1.bold = True
r_t1.font.size = Pt(18)
r_t1.font.color.rgb = NAVY

r_t2 = p_t.add_run('Assurances & Sinistres, Visites & Réforme, Budget & Prévisions, Infractions, Sécurité & Notifications (10/08 – 14/08)')
r_t2.bold = True
r_t2.font.size = Pt(11.5)
r_t2.font.color.rgb = BLUE_ACCENT

p_t.paragraph_format.space_after = Pt(14)

# 1. Informations Générales
h1 = doc.add_heading('1. Informations Générales', level=1)
h1.runs[0].font.color.rgb = NAVY
h1.runs[0].font.size = Pt(12.5)

table_info = doc.add_table(rows=4, cols=2)
table_info.alignment = WD_TABLE_ALIGNMENT.CENTER
table_info.autofit = False

info_data = [
    ('Sprint', 'Sprint 5 (Sprint Final de Clôture du Périmètre Functional CdC)'),
    ('Période prévisionnelle', '10/08/2026 au 14/08/2026 (5 jours ouvrés)'),
    ('Score de Conformité CdC', '100% des exigences fonctionnelles et techniques du Cahier des Charges MEF couvertes.'),
    ('Objectif Principal', 'Finaliser les modules d\'Assurances, Sinistres, Infractions/PV, Visites Techniques, Taxes, Procédure de Réforme des véhicules, Prévisions de consommation de carburant, Suivi Budgétaire Analytique, Sécurité RBAC/Audit, Service d\'Alertes Email (JavaMailSender) et GED Sécurisée.')
]

for idx, (label, val) in enumerate(info_data):
    r_cells = table_info.rows[idx].cells
    r_cells[0].text = label
    r_cells[0].paragraphs[0].runs[0].font.bold = True
    r_cells[0].paragraphs[0].runs[0].font.color.rgb = NAVY
    r_cells[1].text = val

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# 2. Contenu du Sprint 5
h2 = doc.add_heading('2. Contenu du Sprint 5', level=1)
h2.runs[0].font.color.rgb = NAVY
h2.runs[0].font.size = Pt(12.5)

# 2.1 Fonctionnalités Prévues (Backlog Révisé 100% CdC)
h2_1 = doc.add_heading('2.1. Fonctionnalités Prévues (Backlog 100% CdC MEF)', level=2)
h2_1.runs[0].font.color.rgb = BLUE_ACCENT
h2_1.runs[0].font.size = Pt(11)

features = [
    ('1. Assurances, Traitement des Sinistres & Infractions Routières (PV/Amendes)', [
        'Assurances : Suivi des contrats d\'assurance, polices, compagnies (AXA, RMA, Wafa), garanties (tous risques, tiers, vol), primes, franchises et alertes J-30 avant expiration.',
        'Sinistres : Déclaration des accidents (date, lieu, conducteur, tiers impliqués, constats, photos, dommages), expertises d\'assurance, indemnisations et clôture du dossier.',
        'Infractions & Contraventions (CdC Section 4 & 9) : Saisie des PV/amendes radar rattachés aux véhicules et aux conducteurs lors de leurs missions, suivi des paiements et régularisation.'
    ]),
    ('2. Visites Techniques, Taxes Automobiles & Workflow de Réforme des Véhicules', [
        'Visites Techniques & Taxes : Contrôles réglementaires, PV de visite (favorable, contre-visite sous 15j), vignettes automobiles et suivi des statuts (payé, exonéré, en retard).',
        'Procédure de Réforme & Sortie d\'Inventaire (CdC Section 6.4 & 7.1) : Workflow de déclassement des véhicules vétustes/réformés (EN_COURS_DE_REFORME, REFORME, VENDU), attachement du PV de commission de réforme et sortie formelle de l\'inventaire actif du MEF.'
    ]),
    ('3. Suivi Budgétaire Analytique & Moteur de Prévisions Carburant', [
        'Gestion Budgétaire MEF : Allocation annuelle par Direction/Service et nature de dépense (Carburant, Assurance, Entretien, Réparation, Taxes, Visites), suivi Prévu vs Engagé vs Réalisé vs Restant.',
        'Prévisions de Consommation Carburant (CdC Section 10) : Calcul automatique des besoins futurs (Quantité Prévue = km_prévu × conso_moyenne / 100 ; Montant Prévu = Quantité × Prix_prévu) par Direction et par mois.'
    ]),
    ('4. Administration RBAC, Audit Logs, NotificationService (Email) & GED Sécurisée', [
        'Sécurité & Enums Harmonises : Rôles RBAC (Admin, Gestionnaire Central/Local, Financier, Conducteur) et harmonisation de l\'Enum StatutVehicule (DISPONIBLE, AFFECTE, RESERVE, IMMOBILISE, EN_ENTRETIEN, EN_REPARATION, ACCIDENTE, EN_COURS_DE_REFORME, REFORME, VENDU, RESTITUE, ARCHIVE).',
        'Moteur de Notifications Multi-canaux (CdC Section 21) : Configuration de JavaMailSender (SMTP) pour l\'envoi automatique de courriels d\'alerte (Assurance J-30, Contrôle technique, Permis expirant, Dépassement budget).',
        'GED Sécurisée & Audit Logs : Stockage sécurisé sur disque avec validation MIME (PDF, PNG, JPG), prévisualisation native React et journalisation immutable des actions sensibles (AuditLog).'
    ])
]

for title, points in features:
    p_feat = doc.add_paragraph()
    r_f = p_feat.add_run(f'• {title}\n')
    r_f.bold = True
    r_f.font.color.rgb = NAVY
    for pt in points:
        p_pt = doc.add_paragraph()
        p_pt.paragraph_format.left_indent = Inches(0.2)
        r_bullet = p_pt.add_run('- ')
        r_bullet.bold = True
        p_pt.add_run(pt)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# 2.2 Priorisation MoSCoW & Règles de Gestion Strictes
h2_2 = doc.add_heading('2.2. Priorisation MoSCoW & Règles de Gestion Métier', level=2)
h2_2.runs[0].font.color.rgb = BLUE_ACCENT
h2_2.runs[0].font.size = Pt(11)

moscow_text = [
    'Must Have (Impératif) : Assurances, Sinistres, Visites Techniques/Contre-visites, Taxes, Procédure de Réforme, Sécurité RBAC/Audit & Notifications Email (JavaMailSender).',
    'Should Have (Essentiel) : Suivi Budgétaire analytique, Moteur de Prévision Carburant, Gestion des Infractions/PV & GED Sécurisée.',
    'Could Have (Optionnel) : Prévisualisation native des fichiers Word/Excel complexes dans le navigateur web.'
]

for m in moscow_text:
    p_m = doc.add_paragraph()
    p_m.paragraph_format.left_indent = Inches(0.1)
    r_mb = p_m.add_run('📌 ')
    r_mb.font.color.rgb = NAVY
    p_m.add_run(m)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

rules = [
    'RG01 — Une déclaration de sinistre ou le passage en réforme passe automatiquement le véhicule aux statuts ACCIDENTE, EN_REPARATION ou EN_COURS_DE_REFORME et alerte le Gestionnaire Central.',
    'RG02 — Aucun véhicule ne peut être affecté sans contrat d\'assurance actif ; un blocage strict et une notification email SMTP sont émis à J-30 avant expiration.',
    'RG03 — Un résultat de visite technique CONTRE_VISITE_OBLIGATOIRE génère automatiquement un ordre de réparation sous 15 jours.',
    'RG04 — Toute réforme de véhicule exige le téléversement obligatoire du Procès-Verbal de la Commission de Réforme avant validation finale.',
    'RG05 — Toute action de création, modification ou suppression sur des données budgétaires ou administratives est journalisée de façon immutable dans AuditLog.'
]

for rule in rules:
    p_r = doc.add_paragraph()
    p_r.paragraph_format.left_indent = Inches(0.1)
    r_b = p_r.add_run('✔ ')
    r_b.font.color.rgb = BLUE_ACCENT
    p_r.add_run(rule)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# 3. Livrables Attendus
h3 = doc.add_heading('3. Livrables Attendus (Périmètre Final)', level=1)
h3.runs[0].font.color.rgb = NAVY
h3.runs[0].font.size = Pt(12.5)

deliverables = [
    'Base de Données PostgreSQL : Entités JPA Assurance, Sinistre, Infraction, VisiteTechnique, TaxeAutomobile, ReframeVehicule, PrevisionCarburant, BudgetDirection, Utilisateur, Role, DocumentGED et AuditLog avec migration Liquibase/Flyway de Enum StatutVehicule.',
    'Repositories & Services Java : Services métiers incluant NotificationService (JavaMailSender), PrevisionCarburantService, ReframeService, InfractionService, BudgetService, AuditLogService, DocumentGEDService.',
    'Contrôleurs REST & API Swagger UI : Endpoints sécurisés /api/assurances, /api/sinistres, /api/infractions, /api/visites-techniques, /api/reformes, /api/previsions-carburant, /api/budgets, /api/users, /api/notifications et /api/documents.',
    'Frontend React : Vues /assurances, /sinistres, /infractions, /visites-techniques, /reforme, /previsions, /budget, /administration et viewer /documents.',
    'Démonstration & Recette Finale Complète : Vendredi 14 août 2026 à 11h30 (Couverture 100% Cahier des Charges MEF).'
]

for dev in deliverables:
    p_d = doc.add_paragraph()
    p_d.paragraph_format.left_indent = Inches(0.1)
    r_b = p_d.add_run('✔ ')
    r_b.font.color.rgb = NAVY
    p_d.add_run(dev)

output_docx = r'c:\Users\Hassan\Desktop\park auto MEF\Planning_Sprint_5_Assurances_Sinistres_Budget_Securite.docx'
doc.save(output_docx)
print('DOCX Sprint 5 révisé 100% CdC cree avec succes:', output_docx)
